"""L'HTML della narrativa come lo legge una persona (17 ottobre 2026).

La libreria scrive la pagina (``s3dgraphy.exporter.html_exporter``); qui
pyArchInit la completa, senza toccare la libreria:

* **i sottotitoli** — TempluMare porta nella prosa i sottotitoli in Markdown
  («### Past researches»), che la libreria lascia come testo: diventano
  titoli;
* **l'indice** — Enzo, 17 ottobre 2026: «nella narrativa andrebbe messo
  anche un indice». In testa alla pagina, con i capitoli raggruppati per
  parte quando i titoli la dicono («Rilievo e analisi — Introduction»), i
  sottotitoli sotto il loro capitolo e le fonti per ultime.

Funzioni pure su testo HTML: niente Qt, niente database.
"""
from __future__ import annotations

import base64
import html as _html
import io
import json
import os
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

#: «Indice» nelle lingue di pyArchInit
_INDICE = {"it": "Indice", "en": "Contents", "de": "Inhalt", "fr": "Sommaire", "es": "Índice",
           "ca": "Índex", "pt": "Índice", "ro": "Cuprins", "el": "Περιεχόμενα", "ar": "المحتويات"}

#: la parte e il capitolo, nei titoli che la dicono: «Rilievo e analisi — Introduction»
_SEPARATORE = " — "

#: un titolo di capitolo (con o senza attributi: quelli della mappa) o un
#: paragrafo che è un sottotitolo Markdown
_TITOLI = re.compile(r"<h2(\s[^>]*)?>(.*?)</h2>|<p>(#{2,4})\s+(.+?)</p>", re.S)

_CSS = """
/* pyArchInit: l'indice e i sottotitoli */
.pya-indice { border: 1px solid rgba(128,128,128,.3); border-radius: 6px;
  padding: .6rem 1rem .7rem; margin: 1rem 0 2rem; font-size: .92rem;
  font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif; }
.pya-indice .pya-indice-titolo { font-size: 1rem; margin: .2rem 0 .4rem; border: 0; padding: 0; }
.pya-indice ol { margin: .15rem 0 .15rem 1.3rem; padding: 0; }
.pya-indice ol ol { list-style: none; margin-left: .9rem; }
.pya-indice li { margin: .12rem 0; }
.pya-indice a { text-decoration: none; }
.pya-indice a:hover { text-decoration: underline; }
.pya-parte { font-weight: 600; }
h3 { font-size: 1.08rem; margin: 1.5rem 0 .4rem; }
h4 { font-size: 1rem; margin: 1.2rem 0 .3rem; }
"""


def _lingua(lang: Optional[str]) -> str:
    return str(lang or "it")[:2].lower()


def _voci_dei_titoli(testa: str) -> Tuple[str, List[Tuple[int, str, str]]]:
    """Dà un id a ogni capitolo che non l'ha (quelli della mappa lo hanno
    già) e fa titoli dei sottotitoli Markdown; restituisce la pagina e le
    voci dell'indice, (livello, id, testo), nell'ordine della pagina."""
    voci: List[Tuple[int, str, str]] = []
    pezzi: List[str] = []
    fin = 0
    capitolo, sezione = -1, 0
    for m in _TITOLI.finditer(testa):
        pezzi.append(testa[fin:m.start()])
        if m.group(2) is not None:
            capitolo += 1
            sezione = 0
            attributi = m.group(1) or ""
            gia = re.search(r'\bid="([^"]*)"', attributi)
            ident = gia.group(1) if gia else "pya-cap-%d" % capitolo
            pezzi.append(m.group(0) if gia else '<h2 id="%s"%s>%s</h2>' % (ident, attributi, m.group(2)))
            voci.append((1, ident, m.group(2)))
        else:
            tag = "h3" if len(m.group(3)) <= 3 else "h4"
            ident = "pya-sez-%d-%d" % (max(capitolo, 0), sezione)
            sezione += 1
            pezzi.append('<%s id="%s">%s</%s>' % (tag, ident, m.group(4), tag))
            voci.append((2, ident, m.group(4)))
        fin = m.end()
    pezzi.append(testa[fin:])
    return "".join(pezzi), voci


def _parte(titolo: str) -> Tuple[Optional[str], str]:
    if _SEPARATORE in titolo:
        parte, resto = titolo.split(_SEPARATORE, 1)
        if parte.strip() and resto.strip():
            return parte.strip(), resto.strip()
    return None, titolo


def _lista(voci: List[Tuple[int, str, str]], fonti: Optional[Tuple[str, str]]) -> str:
    """Le voci in liste annidate: parte → capitolo → sottotitoli. Una parte
    con un capitolo solo non fa gruppo: il capitolo tiene il titolo intero."""
    capitoli: List[Tuple[str, str, List[Tuple[str, str]]]] = []
    for livello, ident, testo in voci:
        if livello == 1:
            capitoli.append((ident, testo, []))
        elif capitoli:
            capitoli[-1][2].append((ident, testo))

    def voce(ident, testo, sotto):
        figli = "".join('<li><a href="#%s">%s</a></li>' % s for s in sotto)
        return '<li><a href="#%s">%s</a>%s</li>' % (ident, testo, "<ol>%s</ol>" % figli if figli else "")

    righe: List[str] = []
    i = 0
    while i < len(capitoli):
        parte, _resto = _parte(capitoli[i][1])
        j = i
        while j < len(capitoli) and parte is not None and _parte(capitoli[j][1])[0] == parte:
            j += 1
        if parte is not None and j - i >= 2:
            interni = "".join(voce(c[0], _parte(c[1])[1], c[2]) for c in capitoli[i:j])
            righe.append('<li><span class="pya-parte">%s</span><ol>%s</ol></li>' % (parte, interni))
            i = j
        else:
            righe.append(voce(*capitoli[i]))
            i += 1
    if fonti is not None:
        righe.append('<li><a href="#%s">%s</a></li>' % fonti)
    return "<ol>%s</ol>" % "".join(righe)


def add_index(pagina: str, lang: Optional[str] = None) -> str:
    """La pagina con i sottotitoli fatti titoli e l'indice in testa, dopo il
    titolo e la riga dei responsabili. Una pagina che l'ha già resta com'è."""
    if 'class="pya-indice"' in pagina:
        return pagina
    marca = '<section class="em-sources">'
    testa, sep, coda = pagina.partition(marca)
    testa, voci = _voci_dei_titoli(testa)
    fonti = None
    if sep:
        m = re.match(r"\s*<h2>(.*?)</h2>", coda, re.S)
        if m:
            coda = coda[:m.start()] + '<h2 id="pya-fonti">%s</h2>' % m.group(1) + coda[m.end():]
            fonti = ("pya-fonti", m.group(1))
    if not voci:
        return testa + sep + coda
    parola = _INDICE.get(_lingua(lang), _INDICE["en"])
    nav = ('<nav class="pya-indice" id="pya-indice" aria-label="%s">'
           '<h2 class="pya-indice-titolo">%s</h2>%s</nav>\n' % (parola, parola, _lista(voci, fonti)))
    dopo = re.search(r'<p class="em-byline">.*?</p>\s*', testa, re.S) or re.search(r"</h1>\s*", testa)
    if dopo is not None:
        testa = testa[:dopo.end()] + nav + testa[dopo.end():]
    else:
        testa = re.sub(r"(<body[^>]*>)", lambda m: m.group(1) + nav, testa, count=1)
    pagina = testa + sep + coda
    return pagina.replace("</style>", _CSS + "</style>", 1)


# ---------------------------------------------------------------------------
# Le fonti di un capitolo
# ---------------------------------------------------------------------------

@dataclass
class Source:
    """Un documento su cui il capitolo si appoggia, e per che cosa: le
    letture (unità, proprietà, estrattore) che ne vengono."""
    doc_id: str
    name: str
    description: str = ""
    url: str = ""
    uses: List[Tuple[str, str, str]] = field(default_factory=list)


def _uscite(graph) -> Dict[str, Dict[str, List[str]]]:
    indice: Dict[str, Dict[str, List[str]]] = defaultdict(lambda: defaultdict(list))
    for e in getattr(graph, "edges", None) or []:
        indice[e.edge_source][e.edge_type].append(e.edge_target)
    return indice


def _nome(nodo) -> str:
    return str(getattr(nodo, "name", "") or getattr(nodo, "node_id", ""))


def unit_sources(graph, unit_ids) -> List[Source]:
    """I documenti dietro le unità, seguendo la catena dei paradati:
    unità → proprietà (``has_property``) → estrattore, o combinatore che
    unisce più estrattori (``has_data_provenance``, ``combines``) →
    documento (``extracted_from``). Nell'ordine in cui compaiono."""
    uscite = _uscite(graph)
    nodi = {n.node_id: n for n in getattr(graph, "nodes", None) or []}
    fonti: Dict[str, Source] = {}
    for uid in unit_ids:
        unita = nodi.get(uid)
        if unita is None:
            continue
        for pid in uscite[uid]["has_property"]:
            proprieta = nodi.get(pid)
            for prov in uscite[pid]["has_data_provenance"]:
                nodo = nodi.get(prov)
                tipo = getattr(nodo, "node_type", "")
                estrattori = uscite[prov]["combines"] if tipo == "combiner" else [prov]
                for xid in estrattori:
                    for did in uscite[xid]["extracted_from"]:
                        doc = nodi.get(did)
                        if doc is None or getattr(doc, "node_type", "") != "document":
                            continue
                        fonte = fonti.get(did)
                        if fonte is None:
                            fonte = fonti[did] = Source(
                                did, _nome(doc), str(getattr(doc, "description", "") or ""),
                                str(getattr(doc, "url", None) or (getattr(doc, "data", None) or {}).get("url") or ""))
                        uso = (_nome(unita), _nome(proprieta), _nome(nodi.get(xid)))
                        if uso not in fonte.uses:
                            fonte.uses.append(uso)
    return list(fonti.values())


#: gli archi che legano un'unità all'àncora del capitolo
_ARCHI_DELL_ANCORA = ("is_in_activity", "has_first_epoch", "survive_in_epoch")


def chapter_units(graph, chapter) -> List[str]:
    """Le unità di cui parla il capitolo: quelle citate (o i suoi embed,
    come ``chapter_unit_ids``); se non ne ha, quelle dell'attività o
    dell'epoca a cui è ancorato. Un capitolo senza né l'una né l'altra cosa
    — l'introduzione di TempluMare — non ne ha."""
    from s3dgraphy.nodes.stratigraphic_node import StratigraphicNode
    from .em_narrative import chapter_unit_ids
    citate = list(chapter_unit_ids(chapter))
    if citate:
        return citate
    ancora = getattr(chapter, "anchor", None)
    if not ancora:
        return []
    nodi = {n.node_id: n for n in getattr(graph, "nodes", None) or []}
    unita: List[str] = []
    for e in getattr(graph, "edges", None) or []:
        if (e.edge_target == ancora and e.edge_type in _ARCHI_DELL_ANCORA
                and isinstance(nodi.get(e.edge_source), StratigraphicNode)
                and e.edge_source not in unita):
            unita.append(e.edge_source)
    return unita


# ---------------------------------------------------------------------------
# Le fonti come hotspot (Enzo, 17 ottobre 2026: «le fonti doc pdf media come
# fonti devono esserci sempre, e visualizzabili come hotspot sia cliccandoci
# che passandoci con il mouse sopra»)
# ---------------------------------------------------------------------------

_PAROLE = {
    "it": {"chapter_sources": "Fonti del capitolo", "sources": "Fonti", "open": "Apri il file",
           "used_for": "Letture", "close": "Chiudi", "original": "File originale",
           "not_embedded": "Il file non è dentro la pagina (troppo grande): si apre dal percorso "
                           "originale, sul computer che l'ha esportata.",
           "no_file": "Nessun file collegato a questa fonte."},
    "en": {"chapter_sources": "Chapter sources", "sources": "Sources", "open": "Open the file",
           "used_for": "Readings", "close": "Close", "original": "Original file",
           "not_embedded": "The file is not inside the page (too large): it opens from its "
                           "original path, on the computer that exported it.",
           "no_file": "No file linked to this source."},
}


def _parole(lang: Optional[str]) -> Dict[str, str]:
    return _PAROLE.get(_lingua(lang), _PAROLE["en"])


def _chip(doc_id: str, titolo_html: str, classe: str = "") -> str:
    ident = _html.escape(str(doc_id), quote=True)
    return ('<a class="pya-fonte%s" href="#pya-f-%s" role="button" data-src="%s">%s</a>'
            % (classe, ident, ident, titolo_html))


_CSS_FONTI = """
/* pyArchInit: le fonti come hotspot */
[hidden] { display: none !important; }
a.pya-fonte { color: inherit; text-decoration: none; border-bottom: 1px dotted currentColor;
  cursor: pointer; }
.pya-fonti-cap { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
  font-size: .85rem; margin: .8rem 0 1.4rem; color: #555; }
.pya-fonti-cap a.pya-fonte { display: inline-block; border: 1px solid rgba(128,128,128,.45);
  border-radius: 999px; padding: 0 .55rem; margin: .15rem .2rem; }
.em-sources .pya-desc { color: #666; font-size: .92em; }
.pya-tip { position: fixed; z-index: 2000; max-width: 320px; background: #fff; color: #1a1a1a;
  border: 1px solid rgba(0,0,0,.15); border-radius: 8px; box-shadow: 0 8px 24px rgba(0,0,0,.18);
  padding: .6rem .7rem; font: 13px/1.4 -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
  pointer-events: none; }
.pya-tip img { display: block; max-width: 100%; max-height: 220px; margin: 0 0 .4rem; border-radius: 4px; }
.pya-tipo { color: #777; font-size: 11px; text-transform: uppercase; letter-spacing: .05em; margin-left: .4rem; }
.pya-pannello { position: fixed; z-index: 2001; top: 0; right: 0; bottom: 0; width: min(640px, 100vw);
  overflow: auto; background: #fdfdfc; color: #1a1a1a; box-shadow: -8px 0 32px rgba(0,0,0,.25);
  padding: 1rem 1.2rem 2rem; box-sizing: border-box;
  font: 14px/1.5 -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif; }
.pya-pannello h3 { margin: .2rem 2rem .3rem 0; font-size: 1.15rem; }
.pya-chiudi { position: absolute; top: .5rem; right: .7rem; font-size: 1.5rem; border: 0;
  background: none; cursor: pointer; color: inherit; line-height: 1; }
.pya-vista { display: block; width: 100%; height: 72vh; border: 1px solid rgba(128,128,128,.3);
  border-radius: 6px; margin: .6rem 0; background: #fff; }
img.pya-vista { height: auto; }
.pya-percorso { font: 12px/1.4 ui-monospace, Menlo, monospace; color: #666; word-break: break-all; }
.pya-pannello ul { padding-left: 1.2rem; }
@media (prefers-color-scheme: dark) {
  .pya-tip, .pya-pannello { background: #1e2127; color: #e6e6e6; border-color: #2f343d; }
  .pya-fonti-cap, .em-sources .pya-desc, .pya-percorso, .pya-tipo { color: #9aa3ae; }
}
"""

_JS_FONTI = r"""
(function () {
  var F = window.PYA_FONTI || {}, W = window.PYA_PAROLE || {};
  var tip = document.getElementById('pya-tip'), pan = document.getElementById('pya-pannello');
  if (!tip || !pan) { return; }
  function esc(s) { return String(s == null ? '' : s).replace(/[&<>"']/g, function (c) {
    return {'&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'}[c]; }); }
  function blob(f) {
    if (!f.file) { return null; }
    if (!f._url) {
      var bin = atob(f.file), n = bin.length, a = new Uint8Array(n);
      for (var i = 0; i < n; i++) { a[i] = bin.charCodeAt(i); }
      f._url = URL.createObjectURL(new Blob([a], {type: f.mime || 'application/octet-stream'}));
    }
    return f._url;
  }
  function fonte(a) { return a ? F[a.getAttribute('data-src')] : null; }
  function mostra(a) {
    var f = fonte(a); if (!f) { return; }
    tip.innerHTML = (f.preview ? '<img src="' + f.preview + '" alt="">' : '') +
      '<b>' + esc(f.title) + '</b>' + (f.kind ? '<span class="pya-tipo">' + esc(f.kind) + '</span>' : '') +
      (f.description ? '<div>' + esc(f.description) + '</div>' : '');
    tip.hidden = false;
    var r = a.getBoundingClientRect(), w = tip.offsetWidth, h = tip.offsetHeight;
    var x = Math.max(8, Math.min(window.innerWidth - w - 8, r.left));
    var y = r.bottom + 8; if (y + h > window.innerHeight - 8) { y = Math.max(8, r.top - h - 8); }
    tip.style.left = x + 'px'; tip.style.top = y + 'px';
  }
  function apri(a) {
    var f = fonte(a); if (!f) { return; }
    var url = blob(f), vista = '';
    if (url && f.mime === 'application/pdf') { vista = '<iframe class="pya-vista" src="' + url + '" title="' + esc(f.title) + '"></iframe>'; }
    else if (url && /^image\//.test(f.mime || '')) { vista = '<img class="pya-vista" src="' + url + '" alt="' + esc(f.title) + '">'; }
    else if (f.preview) { vista = '<img class="pya-vista" src="' + f.preview + '" alt="' + esc(f.title) + '">'; }
    var usi = (f.uses || []).map(function (u) { return '<li>' + esc(u[0]) + ' — ' + esc(u[1]) + ' <small>(' + esc(u[2]) + ')</small></li>'; }).join('');
    var dove = url || f.link || '';
    pan.innerHTML = '<button type="button" class="pya-chiudi" aria-label="' + esc(W.close) + '">×</button>' +
      '<h3>' + esc(f.title) + (f.kind ? '<span class="pya-tipo">' + esc(f.kind) + '</span>' : '') + '</h3>' +
      (f.description ? '<p>' + esc(f.description) + '</p>' : '') + vista +
      (dove ? '<p><a href="' + esc(dove) + '" target="_blank" rel="noopener">' + esc(W.open) + '</a></p>' :
              (f.path ? '' : '<p>' + esc(W.no_file) + '</p>')) +
      (!url && f.path ? '<p>' + esc(W.not_embedded) + '</p>' : '') +
      (f.path ? '<p class="pya-percorso">' + esc(W.original) + ': ' + esc(f.path) + '</p>' : '') +
      (usi ? '<h4>' + esc(W.used_for) + '</h4><ul>' + usi + '</ul>' : '');
    pan.hidden = false; tip.hidden = true;
    var b = pan.querySelector('.pya-chiudi'); if (b) { b.focus(); }
  }
  document.addEventListener('mouseover', function (e) { var a = e.target.closest && e.target.closest('a.pya-fonte'); if (a) { mostra(a); } });
  document.addEventListener('mouseout', function (e) { var a = e.target.closest && e.target.closest('a.pya-fonte'); if (a) { tip.hidden = true; } });
  document.addEventListener('focusin', function (e) { var a = e.target.closest && e.target.closest('a.pya-fonte'); if (a) { mostra(a); } });
  document.addEventListener('focusout', function () { tip.hidden = true; });
  document.addEventListener('click', function (e) {
    var a = e.target.closest && e.target.closest('a.pya-fonte');
    if (a) { e.preventDefault(); apri(a); return; }
    if (e.target.closest && e.target.closest('.pya-chiudi')) { pan.hidden = true; }
  });
  document.addEventListener('keydown', function (e) { if (e.key === 'Escape') { pan.hidden = true; tip.hidden = true; } });
})();
"""


def _json_script(dati) -> str:
    """JSON da mettere in uno <script>: niente «</» che chiuderebbe il tag."""
    return (json.dumps(dati, ensure_ascii=False).replace("</", "<\\/")
            .replace("\u2028", "\\u2028").replace("\u2029", "\\u2029"))


def _fine_dei_capitoli(pagina: str) -> int:
    """Dove finisce l'ultimo capitolo: prima delle fonti, del piede, della
    mappa o della chiusura della pagina."""
    posti = [pagina.find(m) for m in ('<section class="em-sources">', '<footer class="em-footer">',
                                      '<aside class="pya-map"')]
    posti = [p for p in posti if p >= 0]
    return min(posti) if posti else (pagina.rfind("</body>") if "</body>" in pagina else len(pagina))


def add_sources(pagina: str, per_capitolo: List[List[str]], registro: Dict[str, dict],
                lang: Optional[str] = None) -> str:
    """Le fonti come hotspot: in fondo a ogni capitolo la riga delle sue
    fonti (``per_capitolo``, gli id dei documenti nell'ordine dei capitoli
    della pagina), le schede «fonte» e le menzioni dei documenti, e l'elenco
    delle fonti in fondo alla pagina, completato con quelle dei capitoli.
    Passandoci sopra una scheda con l'anteprima; cliccando, il pannello con
    il file (il PDF o l'immagine, se è dentro la pagina), il percorso
    dell'originale e le letture che ne vengono. ``registro``: id → voce
    (titolo, descrizione, tipo, letture, anteprima, file in base64, mime,
    link, percorso), vedi :func:`sources_registry`."""
    if not registro or 'id="pya-pannello"' in pagina:
        return pagina
    parole = _parole(lang)
    nomi: Dict[str, str] = {}
    for did, voce in registro.items():
        nomi.setdefault(_html.escape(str(voce.get("title") or did)), did)

    def titolo(did):
        return _html.escape(str((registro.get(did) or {}).get("title") or did))

    # 1 · la riga delle fonti in fondo a ogni capitolo
    fine = _fine_dei_capitoli(pagina)
    inizi = [m.start() for m in re.finditer(r"<h2[\s>]", pagina[:fine])]
    for k in range(len(inizi) - 1, -1, -1):
        ids = [d for d in (per_capitolo[k] if k < len(per_capitolo) else []) if d in registro]
        if not ids:
            continue
        riga = ('<div class="pya-fonti-cap"><span>%s:</span> %s</div>\n'
                % (_html.escape(parole["chapter_sources"]), " ".join(_chip(d, titolo(d)) for d in ids)))
        dove = inizi[k + 1] if k + 1 < len(inizi) else fine
        pagina = pagina[:dove] + riga + pagina[dove:]

    # 2 · le schede «fonte» e le menzioni dei documenti
    pagina = re.sub(r'(<div class="em-kind">fonte</div><div class="em-title">)([^<]+)(</div>)',
                    lambda m: m.group(1) + (_chip(nomi[m.group(2)], m.group(2)) if m.group(2) in nomi
                                            else m.group(2)) + m.group(3), pagina)
    pagina = re.sub(r'<em class="em-mention">([^<]+)</em>',
                    lambda m: _chip(nomi[m.group(1)], m.group(1), " em-mention") if m.group(1) in nomi
                    else m.group(0), pagina)

    # 3 · l'elenco delle fonti in fondo: tutte, come hotspot
    tutte: List[str] = []
    for ids in per_capitolo:
        for d in ids:
            if d in registro and d not in tutte:
                tutte.append(d)

    def voce(did):
        desc = str((registro.get(did) or {}).get("description") or "")
        return "<li>%s%s</li>" % (_chip(did, titolo(did)),
                                  '<span class="pya-desc"> — %s</span>' % _html.escape(desc) if desc else "")

    sezione = re.search(r'(<section class="em-sources">.*?<ul>)(.*?)(</ul>)', pagina, re.S)
    if sezione:
        presenti: List[str] = []

        def sostituisci(m):
            nome = m.group(1)
            if nome in nomi:
                presenti.append(nomi[nome])
                return voce(nomi[nome])
            return m.group(0)
        elenco = re.sub(r"<li>([^<]+)</li>", sostituisci, sezione.group(2))
        elenco += "".join(voce(d) for d in tutte if d not in presenti)
        pagina = pagina[:sezione.start()] + sezione.group(1) + elenco + sezione.group(3) + pagina[sezione.end():]
    elif tutte:
        nuova = ('<section class="em-sources"><h2>%s</h2><ul>%s</ul></section>\n'
                 % (_html.escape(parole["sources"]), "".join(voce(d) for d in tutte)))
        dove = _fine_dei_capitoli(pagina)
        pagina = pagina[:dove] + nuova + pagina[dove:]

    # 4 · la scheda al passaggio, il pannello al clic
    usati = {d: registro[d] for d in registro}
    blocco = ('<div id="pya-tip" class="pya-tip" hidden></div>\n'
              '<aside id="pya-pannello" class="pya-pannello" role="dialog" aria-modal="true" hidden></aside>\n'
              "<script>var PYA_FONTI = %s;\nvar PYA_PAROLE = %s;\n%s</script>\n"
              % (_json_script(usati), _json_script(parole), _JS_FONTI))
    chiusura = pagina.rfind("</body>")
    pagina = (pagina[:chiusura] + blocco + pagina[chiusura:]) if chiusura >= 0 else pagina + blocco
    return pagina.replace("</style>", _CSS_FONTI + "</style>", 1)


# ---------------------------------------------------------------------------
# Il registro delle fonti: i file dei documenti, dal database del progetto
# ---------------------------------------------------------------------------

#: oltre questa misura un file non entra nella pagina: si apre dal percorso
MAX_FILE_IN_PAGE = 2 * 1024 * 1024
#: e quando la pagina ne porta già tanti, nemmeno i successivi
MAX_FILES_IN_PAGE = 25 * 1024 * 1024

_MIME = {"pdf": "application/pdf", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
         "gif": "image/gif", "webp": "image/webp", "bmp": "image/bmp", "svg": "image/svg+xml"}
_IMMAGINI = ("png", "jpg", "jpeg", "gif", "webp", "bmp", "tif", "tiff")
_LATO_ANTEPRIMA = 420


def _data_uri(dati: bytes, mime: str) -> str:
    return "data:%s;base64,%s" % (mime, base64.b64encode(dati).decode("ascii"))


def _anteprima_di(percorso: str, ext: str) -> str:
    """La prima pagina del PDF o l'immagine, ridotte, come data URI; vuoto
    quando non si legge."""
    try:
        if ext == "pdf":
            import fitz
            doc = fitz.open(percorso)
            try:
                pagina = doc.load_page(0)
                scala = _LATO_ANTEPRIMA / float(max(pagina.rect.width, pagina.rect.height, 1))
                png = pagina.get_pixmap(matrix=fitz.Matrix(scala, scala), alpha=False).tobytes("png")
            finally:
                doc.close()
            return _data_uri(png, "image/png")
        if ext in _IMMAGINI:
            from PIL import Image
            with Image.open(percorso) as im:
                im = im.convert("RGB")
                im.thumbnail((_LATO_ANTEPRIMA, _LATO_ANTEPRIMA))
                buffer = io.BytesIO()
                im.save(buffer, format="JPEG", quality=82)
            return _data_uri(buffer.getvalue(), "image/jpeg")
    except Exception:                                        # noqa: BLE001
        return ""
    return ""


def sources_registry(graph, fonti: Dict[str, "Source"], connection, thumb_base: str,
                     resize_base: str, *, max_file: int = MAX_FILE_IN_PAGE,
                     budget: int = MAX_FILES_IN_PAGE) -> Dict[str, dict]:
    """Per ogni fonte la voce che la pagina mostra: titolo, descrizione,
    tipo, letture, l'anteprima (prima pagina del PDF, l'immagine, o la
    miniatura dell'importazione), il file dentro la pagina se non è troppo
    grande, il percorso dell'originale e il suo indirizzo ``file://``. I
    file si trovano come quelli della scheda: i media della riga del
    documento (``media_for_unit``), l'originale o la sua copia nella
    cartella delle ridotte. Senza database restano titolo, descrizione e
    letture."""
    from ..utility.em_matrix_links import media_for_unit
    registro: Dict[str, dict] = {}
    speso = 0
    for did, fonte in fonti.items():
        voce = {"title": fonte.name, "description": fonte.description, "kind": "",
                "uses": [list(u) for u in fonte.uses], "preview": "", "file": "", "mime": "",
                "link": "", "path": ""}
        nodo = graph.find_node_by_id(did) if hasattr(graph, "find_node_by_id") else None
        refs = []
        if connection and nodo is not None:
            try:
                refs = media_for_unit(connection, nodo, limit=3)
            except Exception:                                # noqa: BLE001
                refs = []
        if refs:
            ref = refs[0]
            originale = str(ref.original_file or "")
            ext = os.path.splitext(originale or ref.name)[1].lstrip(".").lower()
            copia = (os.path.join(resize_base, "%s_%s.%s" % (ref.id_media, ref.name, ext))
                     if resize_base and ext else "")
            file = (originale if originale and os.path.isfile(originale)
                    else copia if copia and os.path.isfile(copia) else "")
            miniatura = os.path.join(thumb_base, ref.thumb_file) if thumb_base and ref.thumb_file else ""
            voce["kind"] = ext.upper()
            voce["mime"] = _MIME.get(ext, "")
            voce["path"] = originale or file
            if file:
                voce["link"] = Path(os.path.abspath(file)).as_uri()
                voce["preview"] = _anteprima_di(file, ext)
            if not voce["preview"] and miniatura and os.path.isfile(miniatura):
                with open(miniatura, "rb") as fh:
                    voce["preview"] = _data_uri(fh.read(), "image/png")
            if file and voce["mime"]:
                misura = os.path.getsize(file)
                if misura <= max_file and speso + misura <= budget:
                    with open(file, "rb") as fh:
                        voce["file"] = base64.b64encode(fh.read()).decode("ascii")
                    speso += misura
        elif "://" in fonte.url:
            voce["link"] = fonte.url
        registro[did] = voce
    return registro


def _documenti_citati(graph, chapter) -> List[str]:
    """I documenti che il capitolo cita da sé: menzioni e schede."""
    from s3dgraphy.nodes.narrative_node import mentions_in
    ids: List[str] = []
    for b in getattr(chapter, "blocks", None) or []:
        if getattr(b, "block_type", "") == "embed":
            candidati = [getattr(b, "ref", None)]
        else:
            candidati = mentions_in(getattr(b, "text", "") or "")
        for ref in candidati:
            nodo = graph.find_node_by_id(ref) if ref else None
            if nodo is not None and getattr(nodo, "node_type", "") == "document" and ref not in ids:
                ids.append(ref)
    return ids


def sources_for_chapters(graph, narrative) -> Tuple[List[List[str]], Dict[str, Source]]:
    """Le fonti di ogni capitolo (gli id, nell'ordine dei capitoli) e tutte
    le fonti, con le letture di tutti i capitoli: i documenti dietro le
    unità di cui il capitolo parla e quelli che cita da sé."""
    per_capitolo: List[List[str]] = []
    tutte: Dict[str, Source] = {}
    for cap in getattr(narrative, "chapters", None) or []:
        ids: List[str] = []
        for fonte in unit_sources(graph, chapter_units(graph, cap)):
            gia = tutte.get(fonte.doc_id)
            if gia is None:
                tutte[fonte.doc_id] = fonte
            else:
                gia.uses.extend(u for u in fonte.uses if u not in gia.uses)
            ids.append(fonte.doc_id)
        for did in _documenti_citati(graph, cap):
            if did not in tutte:
                doc = graph.find_node_by_id(did)
                tutte[did] = Source(did, _nome(doc), str(getattr(doc, "description", "") or ""),
                                    str(getattr(doc, "url", None)
                                        or (getattr(doc, "data", None) or {}).get("url") or ""))
            if did not in ids:
                ids.append(did)
        per_capitolo.append(ids)
    return per_capitolo, tutte


def align_chapters(pagina: str, titoli: List[str]) -> List[Optional[int]]:
    """Per ogni titolo di capitolo della pagina (prima delle fonti), l'indice
    del capitolo della narrativa: per titolo, nell'ordine; per posizione
    quando il titolo non si ritrova."""
    testa = pagina.split('<section class="em-sources">', 1)[0]
    nella_pagina = [_html.unescape(m.group(2))
                    for m in re.finditer(r"<h2(\s[^>]*)?>(.*?)</h2>", testa, re.S)
                    if "pya-indice-titolo" not in (m.group(1) or "")]
    return align_titles(nella_pagina, titoli)


def align_titles(visti: List[str], attesi: List[str]) -> List[Optional[int]]:
    """Per ogni titolo visto, l'indice del titolo atteso uguale (nell'ordine,
    ognuno una volta), o per posizione quando non c'è."""
    usati: set = set()
    ordine: List[Optional[int]] = []
    for k, titolo in enumerate(visti):
        trovato = next((i for i, t in enumerate(attesi) if t == titolo and i not in usati), None)
        if trovato is None and k < len(attesi) and k not in usati:
            trovato = k
        if trovato is not None:
            usati.add(trovato)
        ordine.append(trovato)
    return ordine


# ---------------------------------------------------------------------------
# Il DOCX: sottotitoli, corsivo, fonti dei capitoli, indice di Word
# ---------------------------------------------------------------------------

_MARCHE = re.compile(r"(\*\*[^*\n]+\*\*|\*[^*\n]+\*)")


def _corsivo_nei_run(paragrafo) -> None:
    """``*x*`` e ``**x**`` dentro ogni run diventano run in corsivo o in
    grassetto, con la formattazione del run da cui vengono (le menzioni
    della libreria restano come sono)."""
    import copy
    from docx.text.run import Run
    for run in list(paragrafo.runs):
        pezzi = [p for p in _MARCHE.split(run.text or "") if p]
        if len(pezzi) < 2 and not (pezzi and _MARCHE.fullmatch(pezzi[0])):
            continue
        precedente = run
        for i, pezzo in enumerate(pezzi):
            corsivo = grassetto = False
            if pezzo.startswith("**") and pezzo.endswith("**") and len(pezzo) > 4:
                pezzo, grassetto = pezzo[2:-2], True
            elif pezzo.startswith("*") and pezzo.endswith("*") and len(pezzo) > 2:
                pezzo, corsivo = pezzo[1:-1], True
            if i == 0:
                corrente = run
            else:
                elemento = copy.deepcopy(run._r)
                precedente._r.addnext(elemento)
                corrente = Run(elemento, paragrafo)
            corrente.text = pezzo
            if corsivo:
                corrente.italic = True
            if grassetto:
                corrente.bold = True
            precedente = corrente


def _spezza_paragrafo(paragrafo) -> list:
    """Il blocco di prosa che la libreria scrive in un paragrafo solo, con
    gli a capo dentro: un paragrafo di Word per ogni paragrafo del testo,
    con il corsivo e il grassetto dei run (le menzioni) conservati."""
    import copy
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    if not re.search(r"\n\s*\n", paragrafo.text or ""):
        return [paragrafo]
    pezzi: List[List[Tuple[str, bool, bool]]] = [[]]
    for run in paragrafo.runs:
        for i, parte in enumerate(re.split(r"\n\s*\n", run.text or "")):
            if i > 0:
                pezzi.append([])
            if parte:
                pezzi[-1].append((parte, bool(run.italic), bool(run.bold)))
    pezzi = [c for c in pezzi if "".join(t for t, _i, _b in c).strip()]
    if len(pezzi) < 2:
        return [paragrafo]
    for run in list(paragrafo.runs):
        run._r.getparent().remove(run._r)
    paragrafi = [paragrafo]
    for _pezzo in pezzi[1:]:
        nuovo = OxmlElement("w:p")
        if paragrafo._p.pPr is not None:
            nuovo.append(copy.deepcopy(paragrafo._p.pPr))
        paragrafi[-1]._p.addnext(nuovo)
        paragrafi.append(Paragraph(nuovo, paragrafo._parent))
    for par, pezzo in zip(paragrafi, pezzi):
        for k, (testo, corsivo, grassetto) in enumerate(pezzo):
            if k == 0:
                testo = testo.lstrip("\n")
            if k == len(pezzo) - 1:
                testo = testo.rstrip("\n")
            run = par.add_run(testo)
            if corsivo:
                run.italic = True
            if grassetto:
                run.bold = True
    return paragrafi


def _campo(tipo: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = OxmlElement("w:r")
    segno = OxmlElement("w:fldChar")
    segno.set(qn("w:fldCharType"), tipo)
    run.append(segno)
    return run


def _istruzione(testo: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = OxmlElement("w:r")
    istruzione = OxmlElement("w:instrText")
    istruzione.set(qn("xml:space"), "preserve")
    istruzione.text = testo
    run.append(istruzione)
    return run


def improve_docx(dati: bytes, graph, narrative, lang: Optional[str] = None) -> bytes:
    """Il DOCX della libreria con i sottotitoli Markdown fatti titoli di
    secondo livello, il corsivo e il grassetto veri, la riga delle fonti in
    fondo a ogni capitolo, le fonti dei capitoli nell'elenco finale e, prima
    dei capitoli, l'indice di Word: un campo TOC con i titoli già scritti,
    che Word aggiorna con i numeri di pagina all'apertura."""
    import docx
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    documento = docx.Document(io.BytesIO(dati))
    parole = _parole(lang)

    # 1 · un paragrafo per paragrafo, poi sottotitoli e corsivo
    paragrafi = []
    for p in list(documento.paragraphs):
        paragrafi.extend(_spezza_paragrafo(p) if p.style.name == "Normal" else [p])
    for p in paragrafi:
        m = re.match(r"^(#{2,4})\s+(.+)$", (p.text or "").strip(), re.S)
        if m and p.style.name == "Normal":
            for run in list(p.runs)[1:]:
                run._r.getparent().remove(run._r)
            if p.runs:
                p.runs[0].text = m.group(2).strip()
            else:
                p.add_run(m.group(2).strip())
            p.style = documento.styles["Heading 2" if len(m.group(1)) <= 3 else "Heading 3"]
        elif "*" in (p.text or ""):
            _corsivo_nei_run(p)

    # 2 · le fonti di ogni capitolo, e quelle che mancano all'elenco finale
    per_capitolo, fonti = sources_for_chapters(graph, narrative)

    def descritta(fonte):
        return fonte.name + (" (%s)" % fonte.description if fonte.description else "")

    capitoli = [p for p in documento.paragraphs if p.style.name == "Heading 1"]
    ordine = align_titles([p.text for p in capitoli],
                          [str(getattr(c, "title", "") or "") for c in narrative.chapters])
    for k, titolo in enumerate(capitoli):
        indice_cap = ordine[k] if k < len(ordine) else None
        ids = per_capitolo[indice_cap] if indice_cap is not None and indice_cap < len(per_capitolo) else []
        if not ids:
            continue
        riga = "%s: %s" % (parole["chapter_sources"], " · ".join(descritta(fonti[d]) for d in ids))
        if k + 1 < len(capitoli):
            capitoli[k + 1].insert_paragraph_before(riga)
        else:
            documento.add_paragraph(riga)
    if fonti:
        presenti = {p.text.strip() for p in documento.paragraphs}
        mancanti = [f for f in fonti.values() if f.name not in presenti]
        intestazioni = {"Fonti", "Sources", parole["sources"]}
        if not any(p.style.name == "Heading 1" and p.text.strip() in intestazioni
                   for p in documento.paragraphs):
            documento.add_paragraph(parole["sources"], style="Heading 1")
            mancanti = list(fonti.values())
        for fonte in mancanti:
            documento.add_paragraph(descritta(fonte))

    # 3 · l'indice, prima del primo capitolo
    primo = next((p for p in documento.paragraphs if p.style.name == "Heading 1"), None)
    if primo is not None:
        voci = [(1 if p.style.name == "Heading 1" else 2, p.text) for p in documento.paragraphs
                if p.style.name in ("Heading 1", "Heading 2")]
        titolo = primo.insert_paragraph_before()
        corsa = titolo.add_run(_INDICE.get(_lingua(lang), _INDICE["en"]))
        corsa.bold = True
        corsa.font.size = Pt(14)
        for i, (livello, testo) in enumerate(voci):
            voce = primo.insert_paragraph_before()
            if livello > 1:
                voce.paragraph_format.left_indent = Pt(14)
            if i == 0:
                voce._p.append(_campo("begin"))
                voce._p.append(_istruzione(' TOC \\o "1-2" \\h \\z \\u '))
                voce._p.append(_campo("separate"))
            voce.add_run(testo)
            if i == len(voci) - 1:
                voce._p.append(_campo("end"))
        primo.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)
        impostazioni = documento.settings.element
        aggiorna = impostazioni.find(qn("w:updateFields"))
        if aggiorna is None:
            aggiorna = OxmlElement("w:updateFields")
            impostazioni.append(aggiorna)
        aggiorna.set(qn("w:val"), "true")

    uscita = io.BytesIO()
    documento.save(uscita)
    return uscita.getvalue()
